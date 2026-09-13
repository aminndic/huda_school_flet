import os
import flet as ft
from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

def main(page: ft.Page):
    page.title = "مدرسة الهدى الثانوية بنين - نظام إدارة النتائج والطلاب"
    page.rtl = True
    page.theme_mode = ft.ThemeMode.LIGHT
    page.padding = 20

    # بيانات سابقة افتراضية للطلاب والنتائج
    students_data = [
        {"id": "101", "name": "أحمد محمد علي", "arabic": 85, "math": 90, "english": 78, "physics": 88},
        {"id": "102", "name": "عمر خالد محمود", "arabic": 92, "math": 95, "english": 89, "physics": 94},
    ]

    # --- 1. تصدير نتيجة طالب محدد إلى ملف PDF ---
    def generate_student_pdf(student):
        downloads_dir = os.path.join(os.path.expanduser("~"), "Downloads")
        if not os.path.exists(downloads_dir):
            os.makedirs(downloads_dir)

        pdf_path = os.path.join(downloads_dir, f"نتيجة_{student['id']}_{student['name']}.pdf")
        
        c = canvas.Canvas(pdf_path, pagesize=A4)
        c.setFont("Helvetica-Bold", 18)
        c.drawString(140, 800, "مدرسة الهدى الثانوية بنين")
        c.setFont("Helvetica-Bold", 14)
        c.drawString(170, 775, "شهادة نتيجة امتحان طالب")
        
        c.setFont("Helvetica", 12)
        c.drawString(50, 720, f"اسم الطالب: {student['name']}")
        c.drawString(50, 700, f"رقم الجلوس / الرقم الأكاديمي: {student['id']}")
        
        c.drawString(50, 660, "المادة")
        c.drawString(250, 660, "الدرجة")
        c.line(50, 650, 450, 650)

        y = 620
        subjects = [
            ("اللغة العربية", student.get("arabic", 0)),
            ("الرياضيات", student.get("math", 0)),
            ("اللغة الإنجليزية", student.get("english", 0)),
            ("الفيزياء", student.get("physics", 0)),
        ]

        total = 0
        for subj, score in subjects:
            c.drawString(50, y, subj)
            c.drawString(250, y, str(score))
            total += int(score)
            y -= 25

        c.line(50, y + 10, 450, y + 10)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, y - 15, f"المجموع الكلي: {total} / 400")
        c.save()

        snack_bar = ft.SnackBar(ft.Text(f"تم تصدير نتيجة الطالب إلى: {pdf_path}"))
        page.overlay.append(snack_bar)
        snack_bar.open = True
        page.update()

    # --- 2. تصدير كامل النتائج إلى ملف Excel باستخدام openpyxl فقط ---
    def export_excel_click(e):
        downloads_dir = os.path.join(os.path.expanduser("~"), "Downloads")
        file_path = os.path.join(downloads_dir, "نتائج_طلاب_مدرسة_الهدى.xlsx")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "النتائج"

        # كتابة العناوين الرئيسية
        headers = ["رقم الجلوس", "اسم الطالب", "اللغة العربية", "الرياضيات", "اللغة الإنجليزية", "الفيزياء"]
        ws.append(headers)

        # كتابة بيانات الطلاب
        for s in students_data:
            ws.append([
                s["id"],
                s["name"],
                s.get("arabic", 0),
                s.get("math", 0),
                s.get("english", 0),
                s.get("physics", 0)
            ])

        wb.save(file_path)
        
        snack_bar = ft.SnackBar(ft.Text(f"تم تصدير ملف Excel بنجاح إلى: {file_path}"))
        page.overlay.append(snack_bar)
        snack_bar.open = True
        page.update()

    # --- 3. استيراد النتائج من ملف Excel باستخدام openpyxl فقط ---
    def on_file_picked(e: ft.FilePickerResultEvent):
        if e.files:
            file_path = e.files[0].path
            try:
                wb = load_workbook(filename=file_path, data_only=True)
                ws = wb.active

                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    return

                # قراءة الهيدر لتحديد مواقع الأعمدة
                header = [str(h).strip() if h is not None else "" for h in rows[0]]
                
                def get_idx(name_list):
                    for name in name_list:
                        if name in header:
                            return header.index(name)
                    return -1

                idx_id = get_idx(["رقم الجلوس", "id", "الرقم"])
                idx_name = get_idx(["اسم الطالب", "name", "الاسم"])
                idx_arabic = get_idx(["اللغة العربية", "arabic", "عربي"])
                idx_math = get_idx(["الرياضيات", "math", "رياضيات"])
                idx_english = get_idx(["اللغة الإنجليزية", "english", "إنجليزي"])
                idx_physics = get_idx(["الفيزياء", "physics", "فيزياء"])

                students_data.clear()

                for row in rows[1:]:
                    if not any(row):
                        continue
                    
                    student_id = str(row[idx_id]) if idx_id != -1 and row[idx_id] is not None else ""
                    student_name = str(row[idx_name]) if idx_name != -1 and row[idx_name] is not None else ""
                    
                    if student_id or student_name:
                        students_data.append({
                            "id": student_id,
                            "name": student_name,
                            "arabic": int(row[idx_arabic]) if idx_arabic != -1 and row[idx_arabic] is not None else 0,
                            "math": int(row[idx_math]) if idx_math != -1 and row[idx_math] is not None else 0,
                            "english": int(row[idx_english]) if idx_english != -1 and row[idx_english] is not None else 0,
                            "physics": int(row[idx_physics]) if idx_physics != -1 and row[idx_physics] is not None else 0,
                        })

                refresh_table()
                snack_bar = ft.SnackBar(ft.Text("تم استيراد قائمة الطلاب والنتائج من ملف Excel بنجاح!"))
                page.overlay.append(snack_bar)
                snack_bar.open = True
                page.update()

            except Exception as ex:
                snack_bar = ft.SnackBar(ft.Text(f"حدث خطأ أثناء استيراد الملف: {str(ex)}"))
                page.overlay.append(snack_bar)
                snack_bar.open = True
                page.update()

    file_picker = ft.FilePicker(on_result=on_file_picked)
    page.overlay.append(file_picker)

    # --- 4. عرض النتائج في جدول تفاعلي ---
    table = ft.DataTable(
        columns=[
            ft.DataColumn(ft.Text("رقم الجلوس")),
            ft.DataColumn(ft.Text("اسم الطالب")),
            ft.DataColumn(ft.Text("عربي")),
            ft.DataColumn(ft.Text("رياضيات")),
            ft.DataColumn(ft.Text("إنجليزي")),
            ft.DataColumn(ft.Text("فيزياء")),
            ft.DataColumn(ft.Text("شهادة PDF")),
        ],
        rows=[]
    )

    def refresh_table():
        table.rows.clear()
        for student in students_data:
            table.rows.append(
                ft.DataRow(
                    cells=[
                        ft.DataCell(ft.Text(student["id"])),
                        ft.DataCell(ft.Text(student["name"])),
                        ft.DataCell(ft.Text(str(student["arabic"]))),
                        ft.DataCell(ft.Text(str(student["math"]))),
                        ft.DataCell(ft.Text(str(student["english"]))),
                        ft.DataCell(ft.Text(str(student["physics"]))),
                        ft.DataCell(
                            ft.IconButton(
                                icon=ft.icons.PICTURE_AS_PDF,
                                tooltip="تصدير شهادة النتيجة PDF",
                                on_click=lambda e, s=student: generate_student_pdf(s)
                            )
                        ),
                    ]
                )
            )
        page.update()

    # --- 5. أدوات إدخال وإضافة طالب يدوي ---
    txt_id = ft.TextField(label="رقم الجلوس", width=120)
    txt_name = ft.TextField(label="اسم الطالب", width=200)
    txt_arabic = ft.TextField(label="عربي", width=80)
    txt_math = ft.TextField(label="رياضيات", width=80)
    txt_english = ft.TextField(label="إنجليزي", width=80)
    txt_physics = ft.TextField(label="فيزياء", width=80)

    def add_student_click(e):
        if txt_id.value and txt_name.value:
            students_data.append({
                "id": txt_id.value,
                "name": txt_name.value,
                "arabic": int(txt_arabic.value or 0),
                "math": int(txt_math.value or 0),
                "english": int(txt_english.value or 0),
                "physics": int(txt_physics.value or 0),
            })
            txt_id.value = ""
            txt_name.value = ""
            txt_arabic.value = ""
            txt_math.value = ""
            txt_english.value = ""
            txt_physics.value = ""
            refresh_table()

    # --- العناصر وأزرار الإجراءات ---
    btn_import = ft.ElevatedButton(
        "استيراد من Excel",
        icon=ft.icons.UPLOAD_FILE,
        on_click=lambda _: file_picker.pick_files(allowed_extensions=["xlsx", "xls"])
    )

    btn_export = ft.ElevatedButton(
        "تصدير الكل إلى Excel",
        icon=ft.icons.DOWNLOAD,
        on_click=export_excel_click
    )

    btn_add = ft.ElevatedButton("إضافة طالب", icon=ft.icons.ADD, on_click=add_student_click)

    header = ft.Row([ft.Text("نظام نتائج الطلاب - مدرسة الهدى الثانوية بنين", size=22, weight=ft.FontWeight.BOLD)], alignment=ft.MainAxisAlignment.CENTER)
    
    input_row = ft.Row([txt_id, txt_name, txt_arabic, txt_math, txt_english, txt_physics, btn_add], alignment=ft.MainAxisAlignment.CENTER, wrap=True)
    
    actions_row = ft.Row([btn_import, btn_export], alignment=ft.MainAxisAlignment.CENTER)

    refresh_table()
    page.add(header, ft.Divider(), input_row, ft.Divider(), actions_row, ft.Divider(), table)

ft.app(target=main)
